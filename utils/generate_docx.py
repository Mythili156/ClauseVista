from docx import Document

doc = Document()
doc.add_heading('VENDOR SERVICE AGREEMENT', 0)

doc.add_paragraph('This VENDOR SERVICE AGREEMENT ("Agreement") is made between:')
doc.add_paragraph('Alpha Corp ("Client") AND Beta Services ("Vendor").')

doc.add_heading('1. SERVICES', level=1)
doc.add_paragraph('The Vendor agrees to provide IT maintenance services as described in Exhibit A.')

doc.add_heading('2. PAYMENT', level=1)
doc.add_paragraph('Client shall pay Vendor INR 50,000 per month upon receipt of invoice. Late payments shall incur interest at 18% per annum.')

doc.add_heading('3. INDEMNIFICATION', level=1)
doc.add_paragraph('Vendor agrees to indemnify Client against any claims arising out of Vendor’s negligence.')

doc.add_heading('4. LIMITATION OF LIABILITY', level=1)
doc.add_paragraph('In no event shall Client be liable for any indirect damages.')

doc.add_heading('5. JURISDICTION', level=1)
doc.add_paragraph('Any disputes shall be subject to the exclusive jurisdiction of the courts in Delhi.')

doc.save('d:/contract bot/sample_data/sample_vendor_agreement.docx')
print("Created sample_vendor_agreement.docx")
